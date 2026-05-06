"""
匹配與替換邏輯
流程：
  1. 先用字典查 abbreviation → full_name（快速命中）
  2. 字典查不到 → 用 (arrival, revenue) 去參考檔配對
  3. 配對成功 → 學習並存入字典
  4. 仍失敗 → 標記為「未匹配」讓同仁手動處理
"""
from __future__ import annotations

import shutil
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

import openpyxl
from openpyxl.styles import Font, PatternFill

import database
from parser import (
    FlatEntry,
    CalendarEntry,
    parse_calendar_report,
)


# 容忍的營收誤差（TWD）- 針對浮點誤差
REVENUE_TOLERANCE = 1.0


@dataclass
class MatchResult:
    """單一資料列的匹配結果"""
    entry: CalendarEntry
    full_name: str | None              # 最終得到的完整名稱
    source: str                        # 'dictionary' / 'reference' / 'unmatched'
    candidates: list[str] = field(default_factory=list)  # 參考檔中符合的候選


@dataclass
class ProcessSummary:
    """整體處理結果摘要"""
    total: int = 0
    from_dictionary: int = 0
    from_reference: int = 0
    newly_learned: int = 0
    unmatched: int = 0
    unchanged: int = 0                  # 沒需要替換（名稱本來就完整）
    unmatched_rows: list[dict] = field(default_factory=list)


# ---------- 匹配 ----------

def _is_likely_abbreviated(name: str) -> bool:
    """判斷一個名稱是否可能被截斷（含 '...' 或結尾 '^'）"""
    if not name:
        return False
    return "..." in name or name.strip().endswith("^")


def _arrival_key(e_arrival, e_rev) -> tuple | None:
    """建立匹配 key，用 (arrival_iso, rounded_rev)"""
    if e_arrival is None or e_rev is None:
        return None
    return (e_arrival.isoformat(), round(float(e_rev), 2))


def _build_reference_index(flat_entries: list[FlatEntry]) -> dict:
    """把扁平參考檔 index 成 {(arrival, guestroom_rev): [FlatEntry, ...]}"""
    idx: dict[tuple, list[FlatEntry]] = {}
    for fe in flat_entries:
        key = _arrival_key(fe.arrival, fe.guestroom_revenue)
        if key is None:
            continue
        idx.setdefault(key, []).append(fe)
    return idx


def _lookup_reference(entry: CalendarEntry, ref_idx: dict,
                      flat_entries: list[FlatEntry]) -> list[FlatEntry]:
    """
    用 (arrival, revenue) 查參考檔，階層式 fallback：
      1. 精確 (arrival, revenue)
      2. 營收誤差容忍
      3. arrival ±2 日 + 營收匹配（Delphi 有時用實際 pickup 日）
      4. 僅營收匹配（需在全資料中唯一，且 > 0）
    """
    if entry.arrival is None or entry.total_revenue is None:
        return []

    target_rev = float(entry.total_revenue)

    # 1. 精確匹配
    key = _arrival_key(entry.arrival, entry.total_revenue)
    if key:
        hits = ref_idx.get(key, [])
        if hits:
            return hits

    # 2. 營收誤差容忍（同日）
    arr_iso = entry.arrival.isoformat()
    tolerant: list[FlatEntry] = []
    for (arr, rev), lst in ref_idx.items():
        if arr == arr_iso and abs(rev - target_rev) <= REVENUE_TOLERANCE:
            tolerant.extend(lst)
    if tolerant:
        return tolerant

    # 3. Arrival ±2 日 + 營收匹配
    if target_rev > 0:
        from datetime import timedelta
        near: list[FlatEntry] = []
        for delta in (-1, 1, -2, 2):
            alt = entry.arrival + timedelta(days=delta)
            for fe in flat_entries:
                if fe.arrival == alt and abs(
                    fe.guestroom_revenue - target_rev
                ) <= REVENUE_TOLERANCE:
                    near.append(fe)
            if near:
                return near

    # 4. 僅營收匹配（若唯一且 > 0）
    if target_rev > 0:
        rev_only = [
            fe for fe in flat_entries
            if abs(fe.guestroom_revenue - target_rev) <= REVENUE_TOLERANCE
        ]
        if len(rev_only) == 1:
            return rev_only

    return []


def match_entries(
    calendar_entries: list[CalendarEntry],
    flat_entries: list[FlatEntry] | None,
) -> tuple[list[MatchResult], ProcessSummary]:
    """
    對日曆檔中的每一筆做匹配。
    - flat_entries 為 None 表示沒上傳參考檔，只用字典
    """
    summary = ProcessSummary(total=len(calendar_entries))
    results: list[MatchResult] = []

    ref_idx = _build_reference_index(flat_entries) if flat_entries else {}

    for entry in calendar_entries:
        raw = entry.raw_name.strip()

        # 若名稱本來就是完整的（沒 ... 沒 ^），仍嘗試從字典更新；否則保留原樣
        looks_abbrev = _is_likely_abbreviated(raw)

        # 1. 字典查
        dict_hit = database.lookup(raw)
        if dict_hit:
            results.append(MatchResult(
                entry=entry, full_name=dict_hit, source="dictionary"
            ))
            summary.from_dictionary += 1
            continue

        # 2. 參考檔查
        ref_hits = _lookup_reference(entry, ref_idx, flat_entries or []) \
            if flat_entries else []
        if ref_hits:
            # 優先採用名稱「不等於」當前 raw 的候選（避免把縮寫再存回縮寫）
            chosen = None
            for fh in ref_hits:
                if fh.full_name != raw:
                    chosen = fh
                    break
            if chosen is None:
                chosen = ref_hits[0]

            results.append(MatchResult(
                entry=entry,
                full_name=chosen.full_name,
                source="reference",
                candidates=[fh.full_name for fh in ref_hits],
            ))

            # 學習：只有「看起來是縮寫」的才寫回字典，避免污染
            if looks_abbrev and chosen.full_name != raw:
                status = database.add_or_update(
                    abbreviation=raw,
                    full_name=chosen.full_name,
                    source="auto",
                )
                if status in ("added", "updated"):
                    summary.newly_learned += 1
            summary.from_reference += 1
            continue

        # 3. 完全沒匹配到
        if not looks_abbrev:
            # 名稱本身不像縮寫 → 視為「不需替換」
            results.append(MatchResult(
                entry=entry, full_name=raw, source="unmatched"
            ))
            summary.unchanged += 1
        else:
            results.append(MatchResult(
                entry=entry, full_name=None, source="unmatched"
            ))
            summary.unmatched += 1
            summary.unmatched_rows.append({
                "sheet": entry.sheet_name,
                "row": entry.row,
                "raw_name": entry.raw_name,
                "arrival": entry.arrival.isoformat() if entry.arrival else None,
                "revenue": entry.total_revenue,
            })

    return results, summary


# ---------- 寫回檔案 ----------

UNMATCHED_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
UNMATCHED_FONT = Font(color="C00000", bold=True)


def apply_replacements(
    source_path: str | Path,
    output_path: str | Path,
    match_results: list[MatchResult],
    highlight_unmatched: bool = True,
) -> None:
    """
    把匹配結果寫入新檔案。保留所有原格式。
    - 成功匹配 → 替換 A 欄文字
    - 未匹配且是縮寫 → A 欄文字不動，加黃底紅字標示
    """
    source_path = Path(source_path)
    output_path = Path(output_path)

    # 先整檔複製（保留所有 style/merged cell/formula），再修改
    shutil.copy2(str(source_path), str(output_path))

    wb = openpyxl.load_workbook(str(output_path))

    # 依 sheet 分組
    by_sheet: dict[str, list[MatchResult]] = {}
    for r in match_results:
        by_sheet.setdefault(r.entry.sheet_name, []).append(r)

    for sheet_name, results in by_sheet.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        for r in results:
            cell = ws.cell(row=r.entry.row, column=1)
            if r.full_name and r.source in ("dictionary", "reference"):
                # 替換文字；保留原 cell 格式
                cell.value = r.full_name
            elif r.source == "unmatched" and r.full_name is None and highlight_unmatched:
                # 標示為待人工處理（不改文字）
                cell.fill = UNMATCHED_FILL
                cell.font = UNMATCHED_FONT

    wb.save(str(output_path))


# ---------- 給外部呼叫的入口 ----------

def process(
    calendar_file: str | Path,
    reference_file: str | Path | None,
    output_file: str | Path,
) -> ProcessSummary:
    """端到端執行：讀兩個檔 → 匹配 → 輸出修正版"""
    database.init_db()

    cal_entries, _metas = parse_calendar_report(calendar_file)

    flat_entries = None
    if reference_file:
        from parser import parse_flat_reference
        flat_entries = parse_flat_reference(reference_file)

    results, summary = match_entries(cal_entries, flat_entries)
    apply_replacements(calendar_file, output_file, results)
    return summary


# ---------- CLI 測試 ----------

if __name__ == "__main__":
    import sys

    if len(sys.argv) < 3:
        print("Usage: python matcher.py <calendar.xlsx> <output.xlsx> [reference.xlsx]")
        sys.exit(1)

    calendar = sys.argv[1]
    output = sys.argv[2]
    reference = sys.argv[3] if len(sys.argv) > 3 else None

    summary = process(calendar, reference, output)
    print(f"\n===== 處理結果 =====")
    print(f"總筆數:         {summary.total}")
    print(f"字典命中:       {summary.from_dictionary}")
    print(f"參考檔匹配:     {summary.from_reference}")
    print(f"新學入字典:     {summary.newly_learned}")
    print(f"無需替換:       {summary.unchanged}")
    print(f"未匹配:         {summary.unmatched}")
    if summary.unmatched_rows:
        print(f"\n未匹配列表:")
        for u in summary.unmatched_rows:
            print(f"  - [{u['sheet']}] R{u['row']} "
                  f"{u['raw_name']} (arr={u['arrival']}, rev={u['revenue']})")
    print(f"\n輸出檔案: {output}")
